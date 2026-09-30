import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_border_box(cell, color="CCCCCC", sz="4"):
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def format_run(run, font_name="Nirmala UI", size_pt=9.5, bold=False, italic=False, color_rgb=(30, 41, 59)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_doc_banner(doc, doc_num, eng_title, hin_title):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.77)
    set_cell_background(cell, "0F172A") # Dark Navy Slate
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    
    r0 = p.add_run(f"MM RIDE  •  DOCUMENT {doc_num}\n")
    format_run(r0, size_pt=9, bold=True, color_rgb=(148, 163, 184))
    
    r1 = p.add_run(eng_title.upper() + "\n")
    format_run(r1, size_pt=12.5, bold=True, color_rgb=(255, 255, 255))
    
    r2 = p.add_run(hin_title)
    format_run(r2, size_pt=11.5, bold=True, color_rgb=(226, 232, 240))
    
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)

def add_bilingual_clause(doc, clause_num, eng_heading, hin_heading, eng_body, hin_body):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    # Heading
    r_eng_h = p.add_run(f"{clause_num}. {eng_heading}\n")
    format_run(r_eng_h, size_pt=10.5, bold=True, color_rgb=(15, 23, 42))
    
    r_hin_h = p.add_run(f"{hin_heading}\n")
    format_run(r_hin_h, size_pt=10, bold=True, color_rgb=(30, 58, 138))
    
    # Body
    p_body = doc.add_paragraph()
    p_body.paragraph_format.space_after = Pt(6)
    p_body.paragraph_format.line_spacing = 1.15
    
    r_eng_b = p_body.add_run(eng_body + "\n")
    format_run(r_eng_b, size_pt=9.5, color_rgb=(30, 41, 59))
    
    r_hin_b = p_body.add_run(hin_body)
    format_run(r_hin_b, size_pt=9, italic=True, color_rgb=(71, 85, 105))

def add_signatures_block(doc, include_witnesses=True, include_thumbs=True):
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(12)
    p_head.paragraph_format.space_after = Pt(6)
    r = p_head.add_run("SIGNATURES & ACKNOWLEDGEMENT / हस्ताक्षर एवं अभिस्वीकृति")
    format_run(r, size_pt=10.5, bold=True, color_rgb=(15, 23, 42))
    
    rows_count = 2 if include_witnesses else 1
    table = doc.add_table(rows=rows_count, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(3.38), Inches(3.38)]
    
    # Row 0: Driver & Fleet Operator
    c_driver = table.cell(0, 0)
    c_op = table.cell(0, 1)
    
    c_driver.width = col_widths[0]
    c_op.width = col_widths[1]
    
    set_cell_border_box(c_driver, "CBD5E1")
    set_cell_border_box(c_op, "CBD5E1")
    set_cell_margins(c_driver, top=100, bottom=100, left=120, right=120)
    set_cell_margins(c_op, top=100, bottom=100, left=120, right=120)
    
    pd = c_driver.paragraphs[0]
    pd.paragraph_format.space_after = Pt(2)
    r = pd.add_run("DRIVER / ड्राइवर:\n")
    format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
    r = pd.add_run("Name: Shivkumar Shankarappa Nindi\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = pd.add_run("DL No: MH14 20100054576\n\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = pd.add_run("Signature: ___________________________\n\n")
    format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
    r = pd.add_run("Date: ______________  Place: ___________")
    format_run(r, size_pt=8.5, color_rgb=(71, 85, 105))
    
    if include_thumbs:
        pd.add_run("\n\nLeft Thumb: [        ]  Right Thumb: [        ]")
    
    po = c_op.paragraphs[0]
    po.paragraph_format.space_after = Pt(2)
    r = po.add_run("FLEET OPERATOR / फ्लीट ऑपरेटर:\n")
    format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
    r = po.add_run("Name: Mohamed Ibrahim\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = po.add_run("Role: Fleet Operator / Vehicle Custodian\n\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = po.add_run("Signature: ___________________________\n\n")
    format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
    r = po.add_run("Date: ______________  Place: ___________")
    format_run(r, size_pt=8.5, color_rgb=(71, 85, 105))

    if include_witnesses:
        c_w1 = table.cell(1, 0)
        c_w2 = table.cell(1, 1)
        
        c_w1.width = col_widths[0]
        c_w2.width = col_widths[1]
        
        set_cell_border_box(c_w1, "CBD5E1")
        set_cell_border_box(c_w2, "CBD5E1")
        set_cell_margins(c_w1, top=100, bottom=100, left=120, right=120)
        set_cell_margins(c_w2, top=100, bottom=100, left=120, right=120)
        
        pw1 = c_w1.paragraphs[0]
        pw1.paragraph_format.space_after = Pt(2)
        r = pw1.add_run("WITNESS 1 / गवाह 1:\n")
        format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
        r = pw1.add_run("Name: _____________________________\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw1.add_run("Address/Mobile: ____________________\n\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw1.add_run("Signature: _________________________\n")
        format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
        r = pw1.add_run("Date: _____________________________")
        format_run(r, size_pt=8.5, color_rgb=(71, 85, 105))
        
        pw2 = c_w2.paragraphs[0]
        pw2.paragraph_format.space_after = Pt(2)
        r = pw2.add_run("WITNESS 2 / गवाह 2:\n")
        format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
        r = pw2.add_run("Name: _____________________________\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw2.add_run("Address/Mobile: ____________________\n\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw2.add_run("Signature: _________________________\n")
        format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
        r = pw2.add_run("Date: _____________________________")
        format_run(r, size_pt=8.5, color_rgb=(71, 85, 105))

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_after = Pt(6)

print("Generator script base framework ready.")
