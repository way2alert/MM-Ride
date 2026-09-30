import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document_pack():
    doc = docx.Document()
    
    # ---------------------------------------------------------
    # Page Setup: A4, 0.65 in margins for maximum printable area
    # ---------------------------------------------------------
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.page_width = Inches(8.27)  # A4 width
    section.page_height = Inches(11.69) # A4 height
    
    # Set default style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Nirmala UI'
    style_normal.font.size = Pt(9)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    
    # Header & Footer setup
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("MM RIDE | DRIVER DOCUMENT PACK  •  VEHICLE: DL9SBH6153 (HERO PASSION PRO)")
    hrun.font.name = "Nirmala UI"
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGBColor(100, 116, 139)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun1 = fp.add_run("MM Ride Driver Record  •  Driver: Shivkumar Shankarappa Nindi  •  ")
    frun1.font.name = "Nirmala UI"
    frun1.font.size = Pt(8)
    frun1.font.color.rgb = RGBColor(148, 163, 184)
    
    # Dynamic Page X of Y fields
    f_p = fp.add_run("Page ")
    f_p.font.name = "Nirmala UI"
    f_p.font.size = Pt(8)
    f_p.font.color.rgb = RGBColor(100, 116, 139)
    fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fld1)
    
    f_of = fp.add_run(" of ")
    f_of.font.name = "Nirmala UI"
    f_of.font.size = Pt(8)
    f_of.font.color.rgb = RGBColor(100, 116, 139)
    fld2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
    fp._p.append(fld2)
    
    # ---------------------------------------------------------
    # Styling Helpers
    # ---------------------------------------------------------
    def set_cell_bg(cell, hex_color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_cell_border_box(cell, color="CCCCCC", sz="4"):
        tcPr = cell._element.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

    def format_run(run, font_name="Nirmala UI", size_pt=9, bold=False, italic=False, color_rgb=(30, 41, 59)):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)
        
        # Crucial for Complex Scripts (Hindi Devanagari) rendering in Word
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rFonts.set(qn('w:eastAsia'), font_name)
        
        if bold:
            bCs = rPr.find(qn('w:bCs'))
            if bCs is None:
                rPr.append(OxmlElement('w:bCs'))
        if italic:
            iCs = rPr.find(qn('w:iCs'))
            if iCs is None:
                rPr.append(OxmlElement('w:iCs'))

    def add_doc_banner(doc_num, eng_title, hin_title):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        cell = table.cell(0, 0)
        cell.width = Inches(6.97)
        set_cell_bg(cell, "0F172A") # Deep Slate Navy
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        
        r0 = p.add_run(f"MM RIDE  •  DOCUMENT {doc_num}\n")
        format_run(r0, size_pt=8.5, bold=True, color_rgb=(148, 163, 184))
        
        r1 = p.add_run(eng_title.upper() + "\n")
        format_run(r1, size_pt=11.5, bold=True, color_rgb=(255, 255, 255))
        
        r2 = p.add_run(hin_title)
        format_run(r2, size_pt=10.5, bold=True, color_rgb=(226, 232, 240))
        
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)

    def add_clause(clause_num, eng_h, hin_h, eng_b, hin_b):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.12
        
        r1 = p.add_run(f"{clause_num}. {eng_h}\n")
        format_run(r1, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
        
        r2 = p.add_run(f"{hin_h}\n")
        format_run(r2, size_pt=9, bold=True, color_rgb=(30, 58, 138))
        
        pb = doc.add_paragraph()
        pb.paragraph_format.space_after = Pt(5)
        pb.paragraph_format.line_spacing = 1.12
        
        r3 = pb.add_run(eng_b + "\n")
        format_run(r3, size_pt=8.5, color_rgb=(30, 41, 59))
        
        r4 = pb.add_run(hin_b)
        format_run(r4, size_pt=8.5, bold=False, italic=False, color_rgb=(51, 65, 85)) # clean, non-italicized Hindi

    def add_signatures_block(include_witnesses=True, include_thumbs=True):
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(8)
        p_head.paragraph_format.space_after = Pt(4)
        r = p_head.add_run("SIGNATURES & ACKNOWLEDGEMENT / हस्ताक्षर एवं अभिस्वीकृति")
        format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
        
        rows_count = 2 if include_witnesses else 1
        table = doc.add_table(rows=rows_count, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        col_widths = [Inches(3.48), Inches(3.48)]
        
        # Row 0: Driver & Fleet Operator
        c_driver = table.cell(0, 0)
        c_op = table.cell(0, 1)
        
        c_driver.width = col_widths[0]
        c_op.width = col_widths[1]
        
        set_cell_border_box(c_driver, "CBD5E1")
        set_cell_border_box(c_op, "CBD5E1")
        set_cell_margins(c_driver, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c_op, top=80, bottom=80, left=100, right=100)
        
        pd = c_driver.paragraphs[0]
        pd.paragraph_format.space_after = Pt(2)
        r = pd.add_run("DRIVER / ड्राइवर:\n")
        format_run(r, size_pt=9, bold=True, color_rgb=(15, 23, 42))
        r = pd.add_run("Name: Shivkumar Shankarappa Nindi\n")
        format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
        r = pd.add_run("DL No: MH14 20100054576\n\n")
        format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
        r = pd.add_run("Signature: ___________________________\n\n")
        format_run(r, size_pt=8.5, bold=True, color_rgb=(30, 41, 59))
        r = pd.add_run("Date: ______________  Place: ___________")
        format_run(r, size_pt=8, color_rgb=(71, 85, 105))
        
        if include_thumbs:
            pd.add_run("\n\nLeft Thumb: [        ]  Right Thumb: [        ]")
        
        po = c_op.paragraphs[0]
        po.paragraph_format.space_after = Pt(2)
        r = po.add_run("FLEET OPERATOR / फ्लीट ऑपरेटर:\n")
        format_run(r, size_pt=9, bold=True, color_rgb=(15, 23, 42))
        r = po.add_run("Name: Mohamed Ibrahim\n")
        format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
        r = po.add_run("Role: Fleet Operator / Vehicle Custodian\n\n")
        format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
        r = po.add_run("Signature: ___________________________\n\n")
        format_run(r, size_pt=8.5, bold=True, color_rgb=(30, 41, 59))
        r = po.add_run("Date: ______________  Place: ___________")
        format_run(r, size_pt=8, color_rgb=(71, 85, 105))

        if include_witnesses:
            c_w1 = table.cell(1, 0)
            c_w2 = table.cell(1, 1)
            
            c_w1.width = col_widths[0]
            c_w2.width = col_widths[1]
            
            set_cell_border_box(c_w1, "CBD5E1")
            set_cell_border_box(c_w2, "CBD5E1")
            set_cell_margins(c_w1, top=80, bottom=80, left=100, right=100)
            set_cell_margins(c_w2, top=80, bottom=80, left=100, right=100)
            
            pw1 = c_w1.paragraphs[0]
            pw1.paragraph_format.space_after = Pt(2)
            r = pw1.add_run("WITNESS 1 / गवाह 1:\n")
            format_run(r, size_pt=9, bold=True, color_rgb=(15, 23, 42))
            r = pw1.add_run("Name: _____________________________\n")
            format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
            r = pw1.add_run("Address/Mobile: ____________________\n\n")
            format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
            r = pw1.add_run("Signature: _________________________\n")
            format_run(r, size_pt=8.5, bold=True, color_rgb=(30, 41, 59))
            r = pw1.add_run("Date: _____________________________")
            format_run(r, size_pt=8, color_rgb=(71, 85, 105))
            
            pw2 = c_w2.paragraphs[0]
            pw2.paragraph_format.space_after = Pt(2)
            r = pw2.add_run("WITNESS 2 / गवाह 2:\n")
            format_run(r, size_pt=9, bold=True, color_rgb=(15, 23, 42))
            r = pw2.add_run("Name: _____________________________\n")
            format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
            r = pw2.add_run("Address/Mobile: ____________________\n\n")
            format_run(r, size_pt=8.5, color_rgb=(30, 41, 59))
            r = pw2.add_run("Signature: _________________________\n")
            format_run(r, size_pt=8.5, bold=True, color_rgb=(30, 41, 59))
            r = pw2.add_run("Date: _____________________________")
            format_run(r, size_pt=8, color_rgb=(71, 85, 105))

        p_end = doc.add_paragraph()
        p_end.paragraph_format.space_after = Pt(4)

    # =========================================================
    # DOCUMENT 1: AGREEMENT
    # =========================================================
    add_doc_banner(1, "Driver–Fleet Work & Vehicle Custody Agreement", "ड्राइवर-फ्लीट कार्य एवं वाहन अभिरक्षा समझौता")
    
    # Party details box
    tbl_p = doc.add_table(rows=3, cols=2)
    tbl_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p.autofit = False
    
    p_data = [
        [("Fleet Operator / Custodian", "Mohamed Ibrahim"), ("Driver Name", "Shivkumar Shankarappa Nindi")],
        [("Driver DOB", "02/01/1985"), ("Driving Licence", "MH14 20100054576")],
        [("Present RC Registered Owner", "Rajesh Saini (RC Transfer: Pending)"), ("Vehicle Registration", "DL9SBH6153 (Hero Passion Pro)")]
    ]
    for r_idx, row in enumerate(p_data):
        for c_idx, (label, val) in enumerate(row):
            cell = tbl_p.cell(r_idx, c_idx)
            cell.width = Inches(3.48)
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_border_box(cell, "E2E8F0")
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{label}: ")
            format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
            r2 = p.add_run(val)
            format_run(r2, size_pt=8, color_rgb=(30, 58, 138) if any(k in val for k in ["Rajesh", "Mohamed", "Shivkumar"]) else (30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    # Clauses
    add_clause(
        1, "Parties & Legal Status", "पक्षकार एवं कानूनी स्थिति",
        "This Agreement is entered into between Mohamed Ibrahim ('Fleet Operator / Vehicle Custodian') and Shivkumar Shankarappa Nindi ('Driver'). The Present Registered Owner of vehicle DL9SBH6153 as per RC is Rajesh Saini. RC transfer is pending and may be completed separately through the applicable legal process. Mohamed Ibrahim acts strictly as the Fleet Operator / Vehicle Custodian and not as the registered RC owner.",
        "यह समझौता मोहम्मद इब्राहिम ('फ्लीट ऑपरेटर / वाहन संरक्षक') और शिवकुमार शंकरप्पा निंदी ('ड्राइवर') के बीच संपन्न हुआ है। वाहन DL9SBH6153 के आरसी के अनुसार वर्तमान पंजीकृत मालिक राजेश सैनी हैं। आरसी हस्तांतरण लंबित है जिसे पृथक कानूनी प्रक्रिया द्वारा पूरा किया जा सकता है। मोहम्मद इब्राहिम केवल फ्लीट ऑपरेटर/वाहन संरक्षक के रूप में कार्य करते हैं, पंजीकृत आरसी मालिक के रूप में नहीं।"
    )
    
    add_clause(
        2, "Vehicle Details & Operational Custody", "वाहन विवरण एवं परिचालन अभिरक्षा",
        "Vehicle Details: Hero MotoCorp Passion Pro 13S-Self-Drum-Cast | Reg No: DL9SBH6153 | Reg Date: 29/03/2017 | Valid Up To: 28/03/2032 | Chassis: MBLHAR181HHB39840 | Engine: HA10ACHHBC2459 | Fuel: Petrol. Driver is granted temporary operational custody of the vehicle solely for authorized MM Ride fleet operations. This agreement does NOT transfer vehicle ownership to the Driver or Fleet Operator.",
        "वाहन विवरण: हीरो मोटोकॉर्प पैशन प्रो | पंजीकरण सं: DL9SBH6153 | पंजीकरण तिथि: 29/03/2017 | वैधता: 28/03/2032 | चेसिस सं: MBLHAR181HHB39840 | इंजन सं: HA10ACHHBC2459 | ईंधन: पेट्रोल। ड्राइवर को केवल अधिकृत MM Ride फ्लीट कार्यों हेतु वाहन की अस्थायी परिचालन अभिरक्षा दी जाती है। यह समझौता ड्राइवर या ऑपरेटर को वाहन का स्वामित्व हस्तांतरित नहीं करता है।"
    )
    
    add_clause(
        3, "Driver Qualification & Eligibility", "ड्राइवर योग्यता एवं पात्रता",
        "Driver warrants that he holds a valid Driving Licence (MH14 20100054576) applicable to the vehicle class, is physically and mentally fit, and possesses valid KYC verification. Driver shall notify the Fleet Operator immediately if his licence is suspended or expired.",
        "ड्राइवर यह पुष्टि करता है कि उसके पास वैध ड्राइविंग लाइसेंस (MH14 20100054576) है, वह शारीरिक व मानसिक रूप से सक्षम है और उसके केवाईसी दस्तावेज सत्यापित हैं। लाइसेंस निलंबित या समाप्त होने पर ड्राइवर तुरंत फ्लीट ऑपरेटर को सूचित करेगा।"
    )
    
    add_clause(
        4, "Duty Hours, Weekly Off & Leave Regulations", "ड्यूटी समय, साप्ताहिक अवकाश एवं छुट्टी नियम",
        "The maximum duty limit is UP TO 12 HOURS PER DAY. This 12-hour duration is a maximum cap for safety and is NOT a compulsory minimum duty requirement. Driver is entitled to 1 weekly off per week subject to operational scheduling. Leave must be requested through the prescribed MM Ride procedure. Emergency leave must be communicated as soon as reasonably possible.",
        "अधिकतम ड्यूटी सीमा प्रति दिन 12 घंटे तक है। यह 12 घंटे की अवधि सुरक्षा हेतु अधिकतम सीमा है, न कि अनिवार्य न्यूनतम ड्यूटी। ड्राइवर को परिचालन अनुसूची के अनुसार प्रति सप्ताह 1 साप्ताहिक अवकाश मिलेगा। छुट्टी का अनुरोध निर्धारित प्रक्रिया द्वारा किया जाएगा और आपातकालीन स्थिति में यथाशीघ्र सूचित किया जाएगा।"
    )

    add_clause(
        5, "Strict Prohibitions & Anti-Subletting", "सख्त प्रतिबंध एवं उप-किराया निषेध",
        "Driver shall NOT: (a) hand over or permit any third person to ride the vehicle; (b) sell, lease, pledge, mortgage, transfer, or sublet the vehicle; (c) use the vehicle for illegal activities or prohibited transport; (d) use another driver's platform account or transfer his own MM Ride/aggregator account to any other person.",
        "ड्राइवर निम्नलिखित नहीं करेगा: (क) किसी अन्य व्यक्ति को वाहन सौंपना या चलाने देना; (ख) वाहन को बेचना, पट्टे पर देना, गिरवी रखना या हस्तांतरित करना; (ग) अवैध गतिविधियों में उपयोग करना; (घ) किसी अन्य व्यक्ति के प्लेटफॉर्म खाते का उपयोग करना या अपना खाता किसी को हस्तांतरित करना।"
    )

    add_clause(
        6, "Vehicle Usage & Personal Use Rules", "वाहन उपयोग एवं व्यक्तिगत प्रयोग नियम",
        "The vehicle is dedicated primarily for authorized MM Ride fleet and aggregator operations. Any personal use by the Driver strictly requires prior written permission from the Fleet Operator. Unauthorized personal usage shall constitute a operational violation.",
        "वाहन प्राथमिक रूप से अधिकृत MM Ride फ्लीट एवं एग्रीगेटर कार्यों के लिए है। ड्राइवर द्वारा किसी भी व्यक्तिगत प्रयोग हेतु फ्लीट ऑपरेटर से पूर्व लिखित अनुमति आवश्यक है। बिना अनुमति व्यक्तिगत उपयोग उल्लंघन माना जाएगा।"
    )

    add_clause(
        7, "Daily Earnings Verification & 50/50 Settlement Formula", "दैनिक आय सत्यापन एवं 50/50 निपटान सूत्र",
        "Daily settlement is based strictly ONLY on VERIFIED earnings finalized by the Fleet Operator/Admin upon physical verification of the platform application/device. Driver cannot self-finalize earnings.\n"
        "Formula:\n"
        "• Net Earnings = Verified Gross Earnings - Applicable Platform Charges\n"
        "• Driver Share = 50% of Net Earnings | Fleet Operator Share = 50% of Net Earnings\n"
        "• Temporary Reserve Hold = 10% of Driver Share\n"
        "• Paid Today = 90% of Driver Share (45% of Net Earnings)\n"
        "Example: Net Earnings ₹2,500 => Driver Share ₹1,250 | Operator Share ₹1,250 | Reserve Hold ₹125 | Paid Today ₹1,125.\n"
        "The 10% reserve is recorded separately for operational security and is NOT an automatic salary deduction. Settlement/adjustment of reserve shall follow written procedure and applicable law. Cash rides are included in verified earnings and cannot be reduced by unverified petrol claims.",
        "दैनिक निपटान केवल फ्लीट ऑपरेटर/एडमिन द्वारा प्लेटफॉर्म ऐप/डिवाइस के भौतिक सत्यापन के बाद सत्यापित आय पर आधारित होगा। ड्राइवर स्वयं आय को अंतिम रूप नहीं दे सकता।\n"
        "गणना सूत्र:\n"
        "• शुद्ध आय = सत्यापित सकल आय - लागू प्लेटफॉर्म शुल्क\n"
        "• ड्राइवर हिस्सा = 50% शुद्ध आय | ऑपरेटर हिस्सा = 50% शुद्ध आय\n"
        "• अस्थायी रिजर्व होल्ड = 10% ड्राइवर हिस्सा\n"
        "• आज देय राशि = 90% ड्राइवर हिस्सा (45% शुद्ध आय)\n"
        "उदाहरण: शुद्ध आय ₹2,500 => ड्राइवर हिस्सा ₹1,250 | ऑपरेटर हिस्सा ₹1,250 | रिजर्व होल्ड ₹125 | आज भुगतान ₹1,125।\n"
        "10% रिजर्व अलग से दर्ज किया जाएगा और यह कोई स्वचालित वेतन कटौती नहीं है। रिजर्व का निपटान लिखित प्रक्रिया और कानून के अनुसार होगा। कैश राइड्स सत्यापित आय में शामिल हैं और ईंधन के नाम पर स्वतः कम नहीं की जा सकतीं।"
    )

    add_clause(
        8, "Fuel / Petrol Rules & Expense Policy", "ईंधन / पेट्रोल नियम एवं खर्च नीति",
        "Petrol/fuel is a separate business expense paid by the Fleet Operator and shall NOT automatically reduce the Driver's 50% earnings share. All fuel entries must record date, time, odometer, fuel level before/after, litres, bill amount, pump details, and bill photo. Fake, duplicate, reused bills, false fuel claims, or claiming ride cash as fuel without verification are strictly prohibited.",
        "पेट्रोल/ईंधन फ्लीट ऑपरेटर का पृथक व्यावसायिक खर्च है और यह ड्राइवर के 50% हिस्से को स्वचालित रूप से कम नहीं करेगा। ईंधन रिकॉर्ड में तारीख, समय, ओडोमीटर, ईंधन स्तर, लीटर, बिल राशि और पेट्रोल पंप विवरण दर्ज करना अनिवार्य है। फर्जी, डुप्लिकेट बिल या बिना सत्यापन कैश राइड का ईंधन में दावा करना सख्त वर्जित है।"
    )

    add_clause(
        9, "GPS & Work-Period Telemetry Monitoring Consent", "GPS एवं कार्य-अवधि निगरानी सहमति",
        "Driver consents to work-period recording of GPS location, duty start/end, route, speed, idle duration, breaks, and device/app status for fleet operations, safety, verification, and fraud prevention. THIS CONSENT RELATES SOLELY TO AUTHORIZED WORK SHIFTS AND DOES NOT AUTHORIZE UNRESTRICTED SURVEILLANCE OF DRIVER'S PERSONAL LIFE. Driver shall not disable GPS, use mock location apps, or tamper with tracking.",
        "ड्राइवर फ्लीट संचालन, सुरक्षा और धोखाधड़ी रोकथाम हेतु केवल ड्यूटी अवधि के दौरान GPS लोकेशन, रूट, गति, आइडल समय, ब्रेक और ऐप स्थिति की रिकॉर्डिंग की अनुमति देता है। यह सहमति केवल अधिकृत कार्य शिफ्टों से संबंधित है और ड्राइवर के व्यक्तिगत जीवन की 24/7 निगरानी की अनुमति नहीं देती है। ड्राइवर GPS बंद या मॉक लोकेशन का उपयोग नहीं करेगा।"
    )

    add_clause(
        10, "Idle Threshold & Break Logging Rules", "निष्क्रिय (आइडल) सीमा एवं ब्रेक नियम",
        "Continuous stationary/idle period exceeding 30 minutes during active duty requires system/log reason entry by the Driver. Official break periods (meals, rest) must be logged with start/end time and duration, and shall be recorded separately from active working time.",
        "ड्यूटी के दौरान 30 मिनट से अधिक लगातार स्थिर/निष्क्रिय (idle) रहने पर ड्राइवर द्वारा कारण दर्ज करना आवश्यक होगा। भोजन/विश्राम के ब्रेक का समय सक्रिय कार्य समय से अलग दर्ज किया जाएगा।"
    )

    add_clause(
        11, "Traffic Law Compliance & Challan Handling", "यातायात नियम अनुपालन एवं चालान प्रबंधन",
        "Driver must observe all traffic rules, wear helmets, and refrain from overspeeding, dangerous riding, mobile phone misuse while riding, or riding under the influence of alcohol/drugs. Traffic challans resulting from driver violations shall be handled according to applicable law and circumstances.",
        "ड्राइवर सभी यातायात नियमों का पालन करेगा, हेलमेट पहनेगा, और ओवरस्पीडिंग, खतरनाक राइडिंग, नशे में गाड़ी चलाने या फोन प्रयोग से दूर रहेगा। ड्राइवर के उल्लंघन से हुए चालान का निपटान कानून और परिस्थितियों के अनुसार होगा।"
    )

    add_clause(
        12, "Accident, Incident & Vehicle Damage Protocol", "दुर्घटना, हादसा एवं वाहन क्षति नियम",
        "In case of an accident or damage: (a) prioritize personal & public safety; (b) contact police/emergency services if required; (c) immediately inform Fleet Operator; (d) preserve evidence and photo proof. Damage assessment shall evaluate previous condition logs, GPS data, inspection, and statements. Automatic damage deductions are prohibited; any financial adjustment shall comply with applicable law and written agreement.",
        "दुर्घटना या क्षति की स्थिति में: (क) सुरक्षा को प्राथमिकता दें; (ख) आवश्यकतानुसार पुलिस/आपात्कालीन सेवा को सूचित करें; (ग) फ्लीट ऑपरेटर को तुरंत सूचित करें; (घ) साक्ष्य सुरक्षित रखें। पूर्व स्थिति, GPS और जांच के आधार पर क्षति का मूल्यांकन होगा। स्वचालित क्षति कटौती प्रतिबंधित है; कोई भी समायोजन कानून और लिखित समझौते के अनुसार होगा।"
    )

    add_clause(
        13, "Termination, Exit & Return of Property", "समाप्ति, निकास एवं संपत्ति वापसी",
        "Upon exit or termination, Driver shall return the vehicle (DL9SBH6153), keys, helmet, GPS equipment, and any issued company property. Final settlement will be prepared after physical vehicle inspection and verification of earnings, fuel, challans, and expenses subject to applicable law.",
        "समाप्ति या निकास पर, ड्राइवर वाहन (DL9SBH6153), चाबियां, हेलमेट, GPS उपकरण और कंपनी की संपत्ति वापस करेगा। अंतिम निपटान भौतिक वाहन निरीक्षण, आय, ईंधन, चालान और खर्चों के सत्यापन के बाद कानून के अनुसार किया जाएगा।"
    )

    add_clause(
        14, "Insurance & Regulatory Compliance Disclaimer", "बीमा एवं कानूनी अनुपालन अस्वीकरण",
        "The vehicle shall be used for passenger/fleet work ONLY to the extent permitted by applicable registration, permit, insurance, transport, and aggregator requirements. Required permissions and insurance coverage must be verified before commencing any activity for which they are required. Commercial insurance coverage is not falsely represented.",
        "वाहन का उपयोग केवल लागू पंजीकरण, परमिट, बीमा, परिवहन एवं एग्रीगेटर नियमों द्वारा अनुमति की सीमा तक ही किया जाएगा। आवश्यक अनुमति एवं बीमा कवरेज का सत्यापन कार्य शुरू करने से पहले किया जाना चाहिए। व्यावसायिक बीमा का कोई झूठा प्रतिनिधित्व नहीं किया गया है।"
    )

    add_clause(
        15, "Electronic Records & Immutable Audit Trail", "इलेक्ट्रॉनिक रिकॉर्ड एवं ऑडिट निष्ठा",
        "Both parties acknowledge that MM Ride system logs, duty timestamps, GPS records, and verified settlement entries serve as authoritative electronic records of operational history.",
        "दोनों पक्ष स्वीकार करते हैं कि MM Ride सिस्टम लॉग, ड्यूटी टाइमस्टैम्प, GPS रिकॉर्ड और सत्यापित निपटान प्रविष्टियां परिचालन इतिहास के प्रामाणिक इलेक्ट्रॉनिक रिकॉर्ड के रूप में कार्य करेंगी।"
    )

    add_clause(
        16, "Governing Law & Severability", "शासी कानून एवं पृथक्करणीयता",
        "This Agreement is governed by the laws of India. If any provision is held invalid, the remaining clauses shall continue in full force and effect.",
        "यह समझौता भारत के कानूनों द्वारा शासित है। यदि कोई प्रावधान अमान्य ठहराया जाता है, तो शेष शर्तें पूरी तरह लागू रहेंगी।"
    )

    add_signatures_block(include_witnesses=True, include_thumbs=True)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 2: KYC CHECKLIST
    # =========================================================
    add_doc_banner(2, "Driver KYC & Verification Checklist", "ड्राइवर केवाईसी एवं सत्यापन चेकलिस्ट")
    
    # Profile Table
    tbl_k1 = doc.add_table(rows=4, cols=2)
    tbl_k1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_k1.autofit = False
    
    k_data = [
        [("Driver Full Name", "Shivkumar Shankarappa Nindi"), ("Date of Birth", "02/01/1985")],
        [("Driving Licence No.", "MH14 20100054576"), ("DL Class / Validity", "MCWG / LMV (Valid)")],
        [("Mobile Number", "_____________________"), ("Current Address", "_________________________________")],
        [("Permanent Address", "_____________________"), ("Verification Date", "____ / ____ / 2026")]
    ]
    for r_idx, row in enumerate(k_data):
        for c_idx, (label, val) in enumerate(row):
            cell = tbl_k1.cell(r_idx, c_idx)
            cell.width = Inches(3.48)
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_border_box(cell, "CBD5E1")
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{label}: ")
            format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
            r2 = p.add_run(val)
            format_run(r2, size_pt=8, color_rgb=(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Verification Items Table
    tbl_k2 = doc.add_table(rows=13, cols=4)
    tbl_k2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_k2.autofit = False
    
    k2_headers = ["S.N.", "Document / Verification Item", "Masked Ref / Doc ID", "Status (Check)"]
    k2_widths = [Inches(0.45), Inches(3.3), Inches(1.9), Inches(1.32)]
    
    hdr_row = tbl_k2.rows[0]
    for c_idx, text in enumerate(k2_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = k2_widths[c_idx]
        set_cell_bg(cell, "1E293B")
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        format_run(r, size_pt=8, bold=True, color_rgb=(255, 255, 255))
        
    k2_items = [
        ("1", "Aadhaar Card (Masked Reference)", "XXXX-XXXX-8452 (Ref)", "[  ] Verified / सत्यापित"),
        ("2", "PAN Card (Masked Reference)", "XXXXX1234X (Ref)", "[  ] Verified / सत्यापित"),
        ("3", "Driving Licence (Front & Back)", "MH14 20100054576", "[  ] Verified / सत्यापित"),
        ("4", "Driver Passport Photo / Photograph", "Attached in File", "[  ] Received / प्राप्त"),
        ("5", "Address Proof (Voter/Utility)", "Ref ID: ___________", "[  ] Verified / सत्यापित"),
        ("6", "Emergency Contact Declaration", "Form Doc 9 Attached", "[  ] Submitted / दर्ज"),
        ("7", "Nominee Details Declaration", "Form Doc 9 Attached", "[  ] Submitted / दर्ज"),
        ("8", "Bank Account / UPI (Masked Ref)", "Bank Ref: XXXXXX", "[  ] Verified / सत्यापित"),
        ("9", "Selfie / Live Photo Verification", "MM Ride App Match", "[  ] Matched / मिलान"),
        ("10", "Original Documents Physical Inspection", "Original DL & Aadhaar", "[  ] Verified / जांचा"),
        ("11", "Police Antecedent / Local Check", "Ref: _______________", "[  ] Pending / Cleared"),
        ("12", "Driver Agreement Signed", "Doc 1 Signed", "[  ] Executed / हस्ताक्षरित")
    ]
    for r_idx, (sn, item, ref, status) in enumerate(k2_items, start=1):
        row_cells = tbl_k2.rows[r_idx].cells
        for c_idx, val in enumerate([sn, item, ref, status]):
            cell = row_cells[c_idx]
            cell.width = k2_widths[c_idx]
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_border_box(cell, "E2E8F0")
            set_cell_margins(cell, top=45, bottom=45, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            format_run(r, size_pt=8, color_rgb=(30, 41, 59))
            
    p_k_note = doc.add_paragraph()
    p_k_note.paragraph_format.space_before = Pt(6)
    p_k_note.paragraph_format.space_after = Pt(6)
    r = p_k_note.add_run("Privacy Protection Notice: Full Aadhaar and bank account numbers are stored securely in encrypted digital vault. Only masked reference tokens appear in printed records.\nगोपनीयता सुरक्षा सूचना: पूर्ण आधार और बैंक खाता संख्या सुरक्षित डिजिटल तिजोरी में संग्रहीत हैं। मुद्रित प्रपत्रों में केवल मुखौटा (masked) संदर्भ आईडी प्रदर्शित होते हैं।")
    format_run(r, size_pt=7.5, color_rgb=(100, 116, 139))
    
    add_signatures_block(include_witnesses=False, include_thumbs=False)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 3: BIKE HANDOVER & CUSTODY FORM
    # =========================================================
    add_doc_banner(3, "Bike Handover & Custody Form", "बाइक सुपुर्दगी एवं अभिरक्षा फॉर्म")
    
    tbl_h1 = doc.add_table(rows=3, cols=2)
    tbl_h1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_h1.autofit = False
    
    h1_data = [
        [("Vehicle Make / Model", "Hero MotoCorp Passion Pro 13S"), ("Registration No.", "DL9SBH6153")],
        [("Chassis No.", "MBLHAR181HHB39840"), ("Engine No.", "HA10ACHHBC2459")],
        [("Handover Date & Time", "____/____/2026 at ____:____ AM/PM"), ("Pickup Depot / Location", "MM Ride Hub Depot")]
    ]
    for r_idx, row in enumerate(h1_data):
        for c_idx, (label, val) in enumerate(row):
            cell = tbl_h1.cell(r_idx, c_idx)
            cell.width = Inches(3.48)
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_border_box(cell, "CBD5E1")
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{label}: ")
            format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
            r2 = p.add_run(val)
            format_run(r2, size_pt=8, color_rgb=(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Inspection Checklist Table
    tbl_h2 = doc.add_table(rows=15, cols=4)
    tbl_h2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_h2.autofit = False
    
    h2_headers = ["Item / Component", "Inspection Status", "Remarks / Scratch / Dent Log", "Verified"]
    h2_widths = [Inches(2.2), Inches(1.5), Inches(2.47), Inches(0.8)]
    
    hdr_row = tbl_h2.rows[0]
    for c_idx, text in enumerate(h2_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = h2_widths[c_idx]
        set_cell_bg(cell, "1E293B")
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [1, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        format_run(r, size_pt=8, bold=True, color_rgb=(255, 255, 255))
        
    h2_items = [
        ("Odometer Reading", "_______ km", "Initial Reading Verified", "[  ] OK"),
        ("Fuel Level in Tank", "____ % / ____ Litres", "Fuel Measured at Pickup", "[  ] OK"),
        ("Front & Rear Tyres", "Tread & Pressure Good", "No cuts/punctures", "[  ] OK"),
        ("Headlight & Tail Light", "Functional / Working", "High/Low Beam OK", "[  ] OK"),
        ("Front & Rear Brakes", "Functional / Working", "Lever & Pedal Firm", "[  ] OK"),
        ("Side Mirrors (Left/Right)", "Both Installed", "Clean & Intact", "[  ] OK"),
        ("Horn & Indicators", "Functional / Working", "All 4 Blinkers OK", "[  ] OK"),
        ("Engine Start / Battery", "Self & Kick Start OK", "Smooth Idling", "[  ] OK"),
        ("Body Panels / Scratches", "Inspected", "Log: __________________", "[  ] OK"),
        ("Ignition & Tank Keys", "2 Keys Handed Over", "Key Ring Included", "[  ] OK"),
        ("Driver Helmet", "1 Helmet Issued", "Visor & Strap Intact", "[  ] OK"),
        ("GPS Tracker Device", "Active & Calibrated", "Device Binding Verified", "[  ] OK"),
        ("Vehicle Documents (RC/Ins)", "Attested Copies", "In Vehicle Pouch", "[  ] OK"),
        ("PUC Certificate", "Valid Copy Provided", "Valid Up To: _________", "[  ] OK")
    ]
    for r_idx, (item, status, remarks, ver) in enumerate(h2_items, start=1):
        row_cells = tbl_h2.rows[r_idx].cells
        for c_idx, val in enumerate([item, status, remarks, ver]):
            cell = row_cells[c_idx]
            cell.width = h2_widths[c_idx]
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_border_box(cell, "E2E8F0")
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [1, 3] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            format_run(r, size_pt=8, color_rgb=(30, 41, 59))
            
    p_ack = doc.add_paragraph()
    p_ack.paragraph_format.space_before = Pt(6)
    p_ack.paragraph_format.space_after = Pt(6)
    r1 = p_ack.add_run("ACKNOWLEDGEMENT STATEMENT / अभिस्वीकृति कथन:\n")
    format_run(r1, size_pt=8.5, bold=True, color_rgb=(15, 23, 42))
    r2 = p_ack.add_run('"I, Shivkumar Shankarappa Nindi, hereby acknowledge receipt of vehicle DL9SBH6153 and all listed accessories in the condition recorded above. I accept operational custody subject to the terms of the Driver–Fleet Work Agreement."\n')
    format_run(r2, size_pt=8, color_rgb=(30, 41, 59))
    r3 = p_ack.add_run('"मैं, शिवकुमार शंकरप्पा निंदी, ऊपर दर्ज स्थिति में वाहन DL9SBH6153 और सभी सूचीबद्ध सामान प्राप्त करने की पुष्टि करता हूं। मैं ड्राइवर-फ्लीट कार्य समझौते की शर्तों के अधीन परिचालन अभिरक्षा स्वीकार करता हूं।"')
    format_run(r3, size_pt=8, color_rgb=(51, 65, 85))
    
    add_signatures_block(include_witnesses=True, include_thumbs=False)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 4: DAILY PICKUP & RETURN RECORD
    # =========================================================
    add_doc_banner(4, "Daily Bike Pickup & Return Record", "दैनिक बाइक पिकअप एवं वापसी रिकॉर्ड")
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(4)
    r = p_desc.add_run("Vehicle: DL9SBH6153  •  Driver: Shivkumar Shankarappa Nindi  •  15 Working Days Register\nवाहन: DL9SBH6153  •  ड्राइवर: शिवकुमार शंकरप्पा निंदी  •  15 कार्य दिवस प्रविष्टि रजिस्टर")
    format_run(r, size_pt=8, color_rgb=(71, 85, 105))
    
    tbl_d4 = doc.add_table(rows=16, cols=12)
    tbl_d4.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_d4.autofit = False
    
    d4_headers = ["Day/Date", "Out Time", "Odo Out", "Fuel Out", "In Time", "Odo In", "Fuel In", "Km Run", "Condition", "Key", "Driver Sig", "Op Sig"]
    d4_widths = [Inches(0.7), Inches(0.55), Inches(0.55), Inches(0.5), Inches(0.55), Inches(0.55), Inches(0.5), Inches(0.5), Inches(0.65), Inches(0.35), Inches(0.57), Inches(0.5)]
    
    hdr_row = tbl_d4.rows[0]
    for c_idx, text in enumerate(d4_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = d4_widths[c_idx]
        set_cell_bg(cell, "1E293B")
        set_cell_margins(cell, top=50, bottom=50, left=25, right=25)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        format_run(r, size_pt=7, bold=True, color_rgb=(255, 255, 255))
        
    for r_idx in range(1, 16):
        row_cells = tbl_d4.rows[r_idx].cells
        for c_idx in range(12):
            cell = row_cells[c_idx]
            cell.width = d4_widths[c_idx]
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_border_box(cell, "CBD5E1")
            set_cell_margins(cell, top=55, bottom=55, left=25, right=25)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c_idx == 0:
                r = p.add_run(f"Day {r_idx}")
                format_run(r, size_pt=7, color_rgb=(100, 116, 139))
                
    p_d4_foot = doc.add_paragraph()
    p_d4_foot.paragraph_format.space_before = Pt(6)
    p_d4_foot.paragraph_format.space_after = Pt(6)
    r = p_d4_foot.add_run("Monotonicity Rule: Return Odometer must be greater than or equal to Pickup Odometer. Keys & helmet must be inspected daily.\nओडोमीटर नियम: वापसी ओडोमीटर पिकअप ओडोमीटर से कम नहीं हो सकता। चाबी व हेलमेट की रोजाना जांच अनिवार्य है।")
    format_run(r, size_pt=7.5, color_rgb=(100, 116, 139))
    
    add_signatures_block(include_witnesses=False, include_thumbs=False)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 5: EARNINGS VERIFICATION & SETTLEMENT RECORD
    # =========================================================
    add_doc_banner(5, "Earnings Verification & Settlement Record", "आय सत्यापन एवं भुगतान निपटान रिकॉर्ड")
    
    # Formula Box
    tbl_f = doc.add_table(rows=1, cols=1)
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_f.autofit = False
    cell_f = tbl_f.cell(0, 0)
    cell_f.width = Inches(6.97)
    set_cell_bg(cell_f, "EFF6FF")
    set_cell_border_box(cell_f, "93C5FD")
    set_cell_margins(cell_f, top=80, bottom=80, left=100, right=100)
    pf = cell_f.paragraphs[0]
    pf.paragraph_format.space_after = Pt(0)
    r = pf.add_run("SETTLEMENT FORMULA SUMMARY / निपटान सूत्र विवरण:\n")
    format_run(r, size_pt=8.5, bold=True, color_rgb=(30, 58, 138))
    r = pf.add_run("1. Net Earnings = Verified Gross Earnings - Applicable Platform Charges\n"
                   "2. Driver Share = 50% of Net Earnings  |  Fleet Operator Share = 50% of Net Earnings\n"
                   "3. Temporary Reserve Hold = 10% of Driver Share  |  Paid Today = 90% of Driver Share (45% of Net)\n"
                   "Example: Net ₹2,500 => Driver Share ₹1,250 | Operator Share ₹1,250 | Reserve Hold ₹125 | Paid Today ₹1,125\n"
                   "शुद्ध आय = सत्यापित सकल आय - प्लेटफॉर्म शुल्क | ड्राइवर = 50% | ऑपरेटर = 50% | रिजर्व = 10% ड्राइवर हिस्सा | आज भुगतान = 90% ड्राइवर हिस्सा")
    format_run(r, size_pt=8, color_rgb=(30, 41, 59))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Settlement Table
    tbl_s = doc.add_table(rows=14, cols=10)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_s.autofit = False
    
    s_headers = ["Date", "Platform", "Gross (₹)", "Charges (₹)", "Net (₹)", "Drv 50% (₹)", "Res 10% (₹)", "Paid Today", "Op 50% (₹)", "Txn / Sig"]
    s_widths = [Inches(0.65), Inches(0.72), Inches(0.67), Inches(0.67), Inches(0.67), Inches(0.77), Inches(0.72), Inches(0.77), Inches(0.77), Inches(0.56)]
    
    hdr_row = tbl_s.rows[0]
    for c_idx, text in enumerate(s_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = s_widths[c_idx]
        set_cell_bg(cell, "1E293B")
        set_cell_margins(cell, top=50, bottom=50, left=25, right=25)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        format_run(r, size_pt=7, bold=True, color_rgb=(255, 255, 255))
        
    for r_idx in range(1, 14):
        row_cells = tbl_s.rows[r_idx].cells
        for c_idx in range(10):
            cell = row_cells[c_idx]
            cell.width = s_widths[c_idx]
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_border_box(cell, "CBD5E1")
            set_cell_margins(cell, top=50, bottom=50, left=25, right=25)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c_idx == 0:
                r = p.add_run(f"Entry {r_idx}")
                format_run(r, size_pt=7, color_rgb=(100, 116, 139))
                
    p_s_dec = doc.add_paragraph()
    p_s_dec.paragraph_format.space_before = Pt(6)
    p_s_dec.paragraph_format.space_after = Pt(6)
    r1 = p_s_dec.add_run("OPERATOR VERIFICATION DECLARATION / ऑपरेटर सत्यापन घोषणा:\n")
    format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
    r2 = p_s_dec.add_run('"I confirm that earnings recorded above are physically verified from driver platform app records. Cash rides are included. Reserve is held separately per rules and is not an automatic salary deduction."\n')
    format_run(r2, size_pt=7.5, color_rgb=(30, 41, 59))
    r3 = p_s_dec.add_run('"मैं पुष्टि करता हूं कि ऊपर दर्ज आय का प्लेटफॉर्म ऐप से भौतिक सत्यापन किया गया है। कैश राइड शामिल हैं। रिजर्व नियमों के अनुसार पृथक रखा गया है।"')
    format_run(r3, size_pt=7.5, color_rgb=(51, 65, 85))
    
    add_signatures_block(include_witnesses=False, include_thumbs=False)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 6: FUEL / PETROL RECORD
    # =========================================================
    add_doc_banner(6, "Fuel / Petrol Record & Expense Audit Log", "ईंधन / पेट्रोल रिकॉर्ड एवं खर्च लेखा परीक्षा लॉग")
    
    p_f_intro = doc.add_paragraph()
    p_f_intro.paragraph_format.space_after = Pt(4)
    r = p_f_intro.add_run("Fuel Policy Note: Petrol is a Fleet Operator business expense and shall not automatically reduce Driver 50% share.\nईंधन नीति: पेट्रोल फ्लीट ऑपरेटर का व्यावसायिक खर्च है और यह ड्राइवर के 50% हिस्से को कम नहीं करेगा।")
    format_run(r, size_pt=8, color_rgb=(30, 58, 138))
    
    tbl_f6 = doc.add_table(rows=12, cols=9)
    tbl_f6.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_f6.autofit = False
    
    f6_headers = ["Date & Time", "Vehicle No.", "Odometer", "Fuel Pre", "Litres", "Amount (₹)", "Fuel Post", "Pump & Bill No.", "Verified By"]
    f6_widths = [Inches(0.9), Inches(0.77), Inches(0.67), Inches(0.57), Inches(0.57), Inches(0.72), Inches(0.57), Inches(1.25), Inches(0.95)]
    
    hdr_row = tbl_f6.rows[0]
    for c_idx, text in enumerate(f6_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = f6_widths[c_idx]
        set_cell_bg(cell, "1E293B")
        set_cell_margins(cell, top=50, bottom=50, left=25, right=25)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        format_run(r, size_pt=7, bold=True, color_rgb=(255, 255, 255))
        
    for r_idx in range(1, 12):
        row_cells = tbl_f6.rows[r_idx].cells
        for c_idx in range(9):
            cell = row_cells[c_idx]
            cell.width = f6_widths[c_idx]
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_border_box(cell, "CBD5E1")
            set_cell_margins(cell, top=50, bottom=50, left=25, right=25)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c_idx == 0:
                r = p.add_run(f"Fill {r_idx}")
                format_run(r, size_pt=7, color_rgb=(100, 116, 139))
            elif c_idx == 1:
                r = p.add_run("DL9SBH6153")
                format_run(r, size_pt=7, color_rgb=(30, 41, 59))
                
    p_f_dec = doc.add_paragraph()
    p_f_dec.paragraph_format.space_before = Pt(6)
    p_f_dec.paragraph_format.space_after = Pt(6)
    r1 = p_f_dec.add_run("ANTI-FRAUD FUEL DECLARATION / ईंधन जालसाजी विरोधी घोषणा:\n")
    format_run(r1, size_pt=8, bold=True, color_rgb=(185, 28, 28))
    r2 = p_f_dec.add_run('"Fake bills, duplicate fuel receipts, claiming fuel for unapproved vehicles, or deducting unverified fuel from ride cash is strictly prohibited and subject to fraud investigation."\n')
    format_run(r2, size_pt=7.5, color_rgb=(30, 41, 59))
    r3 = p_f_dec.add_run('"फर्जी बिल, डुप्लिकेट रसीदें या बिना सत्यापन कैश राइड में से ईंधन दावा करना सख्त मना है और धोखाधड़ी जांच के अधीन होगा।"')
    format_run(r3, size_pt=7.5, color_rgb=(51, 65, 85))
    
    add_signatures_block(include_witnesses=False, include_thumbs=False)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 7: GPS / WORK MONITORING CONSENT
    # =========================================================
    add_doc_banner(7, "GPS / Device / Work Monitoring Consent", "GPS / डिवाइस / कार्य निगरानी सहमति")
    
    add_clause(
        1, "Work Telemetry & Location Consent", "कार्य टेलीमेट्री एवं लोकेशन सहमति",
        "Driver consents to MM Ride recording work-time GPS location, duty start/end timestamps, trip routes, speed, idle duration, breaks, and MM Ride driver app status for fleet operations, safety, and fraud prevention.",
        "ड्राइवर फ्लीट संचालन, सुरक्षा और धोखाधड़ी रोकथाम हेतु कार्य-समय के GPS लोकेशन, ड्यूटी समय, ट्रिप रूट, गति, आइडल समय और ऐप स्थिति की रिकॉर्डिंग की अनुमति देता है।"
    )
    
    add_clause(
        2, "Strict Work-Period Scope Disclaimer", "सख्त कार्य-अवधि दायरा अस्वीकरण",
        "THIS CONSENT RELATES SOLELY TO AUTHORIZED WORK SHIFTS / FLEET OPERATIONS AND DOES NOT AUTHORIZE UNRESTRICTED MONITORING OF THE DRIVER'S PERSONAL LIFE OUTSIDE DUTY HOURS.",
        "यह सहमति केवल अधिकृत कार्य शिफ्टों / फ्लीट संचालन से संबंधित है और ड्यूटी घंटों के बाहर ड्राइवर के व्यक्तिगत जीवन की अनियंत्रित निगरानी की अनुमति नहीं देती है।"
    )

    add_clause(
        3, "Anti-Tampering & Device Binding", "छेड़छाड़ निषेध एवं डिवाइस बाइंडिंग",
        "Driver agrees not to disable GPS tracking during shifts, use mock location software, tamper with telemetry hardware/app, or share his login device with unauthorized persons.",
        "ड्राइवर सहमत है कि वह शिफ्ट के दौरान GPS बंद नहीं करेगा, मॉक लोकेशन का उपयोग नहीं करेगा, उपकरण/ऐप से छेड़छाड़ नहीं करेगा या किसी अनधिकृत व्यक्ति से डिवाइस साझा नहीं करेगा।"
    )
    
    add_signatures_block(include_witnesses=False, include_thumbs=False)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 8: ACCIDENT / DAMAGE / CHALLAN REPORT
    # =========================================================
    add_doc_banner(8, "Accident / Damage / Challan Incident Report", "दुर्घटना / क्षति / चालान हादसा रिपोर्ट")
    
    tbl_a8 = doc.add_table(rows=6, cols=2)
    tbl_a8.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_a8.autofit = False
    
    a8_data = [
        [("Incident Date & Time", "____/____/2026 at ____:____ AM/PM"), ("Incident Location", "_________________________________")],
        [("Vehicle Registration", "DL9SBH6153 (Hero Passion Pro)"), ("Driver Name", "Shivkumar Shankarappa Nindi")],
        [("Incident Category", "[  ] Accident  [  ] Damage  [  ] Challan  [  ] Theft"), ("Police Station & GD/FIR No.", "_________________________________")],
        [("Injuries Reported", "[  ] None  [  ] Minor  [  ] Hospitalized"), ("Insurance Intimated", "[  ] Yes  [  ] No  (Claim Ref: ________)")],
        [("Detailed Incident Description", "___________________________________________________________________________________\n___________________________________________________________________________________"), ("Witness Name & Mobile", "_________________________________\nMob: ____________________________")],
        [("Initial Damage Assessment", "___________________________________________________________________________________"), ("Action / Resolution", "___________________________________________________________________________________")]
    ]
    for r_idx, row in enumerate(a8_data):
        for c_idx, (label, val) in enumerate(row):
            cell = tbl_a8.cell(r_idx, c_idx)
            cell.width = Inches(3.48)
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_border_box(cell, "CBD5E1")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{label}:\n")
            format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
            r2 = p.add_run(val)
            format_run(r2, size_pt=8, color_rgb=(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    p_a8_note = doc.add_paragraph()
    p_a8_note.paragraph_format.space_after = Pt(4)
    r1 = p_a8_note.add_run("Legal Note: Damage/Challan liability shall be determined based on evidence and applicable law. Automatic deductions are prohibited.\nकानूनी टिप्पणी: क्षति/चालान देनदारी साक्ष्य और लागू कानून के अनुसार निर्धारित होगी। स्वचालित कटौती वर्जित है।")
    format_run(r1, size_pt=7.5, color_rgb=(100, 116, 139))
    
    add_signatures_block(include_witnesses=True, include_thumbs=False)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 9: EMERGENCY CONTACT & NOMINEE
    # =========================================================
    add_doc_banner(9, "Emergency Contact & Nominee Declaration", "आपातकालीन संपर्क एवं नामांकित व्यक्ति घोषणा")
    
    tbl_e9 = doc.add_table(rows=1, cols=2)
    tbl_e9.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_e9.autofit = False
    
    c_em = tbl_e9.cell(0, 0)
    c_nom = tbl_e9.cell(0, 1)
    
    c_em.width = Inches(3.48)
    c_nom.width = Inches(3.48)
    set_cell_border_box(c_em, "CBD5E1")
    set_cell_border_box(c_nom, "CBD5E1")
    set_cell_margins(c_em, top=80, bottom=80, left=100, right=100)
    set_cell_margins(c_nom, top=80, bottom=80, left=100, right=100)
    
    p = c_em.paragraphs[0]
    r = p.add_run("EMERGENCY CONTACT / आपातकालीन संपर्क:\n")
    format_run(r, size_pt=9, bold=True, color_rgb=(15, 23, 42))
    r = p.add_run("Full Name: _________________________________\n"
                  "Relationship: _____________________________\n"
                  "Primary Mobile: ___________________________\n"
                  "Alternate Mobile: _________________________\n"
                  "Residential Address: ______________________\n"
                  "__________________________________________")
    format_run(r, size_pt=8, color_rgb=(30, 41, 59))
    
    p = c_nom.paragraphs[0]
    r = p.add_run("NOMINEE DECLARATION / नामांकित व्यक्ति:\n")
    format_run(r, size_pt=9, bold=True, color_rgb=(15, 23, 42))
    r = p.add_run("Full Name: _________________________________\n"
                  "Relationship: _____________________________\n"
                  "Mobile Number: ____________________________\n"
                  "Aadhaar Ref: ______________________________\n"
                  "Residential Address: ______________________\n"
                  "__________________________________________")
    format_run(r, size_pt=8, color_rgb=(30, 41, 59))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    p_e9_dec = doc.add_paragraph()
    p_e9_dec.paragraph_format.space_after = Pt(6)
    r1 = p_e9_dec.add_run("DECLARATION & LEGAL DISCLAIMER / घोषणा एवं कानूनी अस्वीकरण:\n")
    format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
    r2 = p_e9_dec.add_run('"I declare that the above emergency and nominee details are true and provided voluntarily. This record is maintained for operational emergency communication and business records, and does not by itself override applicable legal succession laws."\n')
    format_run(r2, size_pt=7.5, color_rgb=(30, 41, 59))
    r3 = p_e9_dec.add_run('"मैं घोषणा करता हूं कि आपातकालीन और नामांकित विवरण सत्य और स्वेच्छा से प्रदान किए गए हैं। यह रिकॉर्ड आपात्कालीन संपर्क हेतु है और उत्तराधिकार कानूनों का उल्लंघन नहीं करता है।"')
    format_run(r3, size_pt=7.5, color_rgb=(51, 65, 85))
    
    add_signatures_block(include_witnesses=False, include_thumbs=True)
    doc.add_page_break()

    # =========================================================
    # DOCUMENT 10: FINAL BIKE RETURN / EXIT FORM
    # =========================================================
    add_doc_banner(10, "Final Bike Return & Exit Clearance Form", "अंतिम बाइक वापसी एवं निकास अनापत्ति फॉर्म")
    
    tbl_x10 = doc.add_table(rows=3, cols=2)
    tbl_x10.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_x10.autofit = False
    
    x_data = [
        [("Driver Name", "Shivkumar Shankarappa Nindi"), ("Vehicle Registration", "DL9SBH6153 (Hero Passion Pro)")],
        [("Exit Date & Time", "____/____/2026 at ____:____ AM/PM"), ("Final Odometer Reading", "___________ km")],
        [("Final Fuel Level", "____ % / ____ Litres"), ("Return Location / Depot", "MM Ride Hub Depot")]
    ]
    for r_idx, row in enumerate(x_data):
        for c_idx, (label, val) in enumerate(row):
            cell = tbl_x10.cell(r_idx, c_idx)
            cell.width = Inches(3.48)
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_border_box(cell, "CBD5E1")
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{label}: ")
            format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
            r2 = p.add_run(val)
            format_run(r2, size_pt=8, color_rgb=(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Return Checklist Table
    tbl_x2 = doc.add_table(rows=10, cols=3)
    tbl_x2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_x2.autofit = False
    
    x2_headers = ["Item / Property Description", "Return Status", "Verification & Remarks"]
    x2_widths = [Inches(2.9), Inches(1.8), Inches(2.27)]
    
    hdr_row = tbl_x2.rows[0]
    for c_idx, text in enumerate(x2_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = x2_widths[c_idx]
        set_cell_bg(cell, "1E293B")
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        format_run(r, size_pt=8, bold=True, color_rgb=(255, 255, 255))
        
    x2_items = [
        ("Vehicle (DL9SBH6153) Handed Back", "[  ] Received", "Body & Frame Inspected"),
        ("Ignition & Tank Keys (2 Keys)", "[  ] Received", "Keys Verified"),
        ("Driver Helmet Handed Back", "[  ] Received", "Condition: ___________"),
        ("GPS Tracker Equipment", "[  ] Received / Bound", "Hardware Intact"),
        ("Company Device / Mobile (if issued)", "[  ] N/A  [  ] Received", "Data Cleared"),
        ("Pending Daily Settlements", "[  ] Cleared  [  ] Pending", "Final Audit Done"),
        ("Temporary Reserve Account Status", "[  ] Settled / Refunded", "Reserve Audited"),
        ("Pending Traffic Challans Check", "[  ] No Pending Challan", "Portal Checked"),
        ("Final Exit Clearance Granted", "[  ] APPROVED", "Account Closed")
    ]
    for r_idx, (item, status, rem) in enumerate(x2_items, start=1):
        row_cells = tbl_x2.rows[r_idx].cells
        for c_idx, val in enumerate([item, status, rem]):
            cell = row_cells[c_idx]
            cell.width = x2_widths[c_idx]
            set_cell_bg(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_border_box(cell, "E2E8F0")
            set_cell_margins(cell, top=45, bottom=45, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            format_run(r, size_pt=8, color_rgb=(30, 41, 59))
            
    p_x_dec = doc.add_paragraph()
    p_x_dec.paragraph_format.space_before = Pt(6)
    p_x_dec.paragraph_format.space_after = Pt(6)
    r1 = p_x_dec.add_run("FINAL CLEARANCE ACKNOWLEDGEMENT / अंतिम अनापत्ति अभिस्वीकृति:\n")
    format_run(r1, size_pt=8, bold=True, color_rgb=(15, 23, 42))
    r2 = p_x_dec.add_run('"Both Driver and Fleet Operator confirm that vehicle DL9SBH6153 and all company properties have been returned and accounts finalized subject to applicable law."\n')
    format_run(r2, size_pt=7.5, color_rgb=(30, 41, 59))
    r3 = p_x_dec.add_run('"दोनों पक्ष पुष्टि करते हैं कि वाहन DL9SBH6153 और सभी संपत्तियां वापस कर दी गई हैं तथा खाते अंतिम रूप से निपटा दिए गए हैं।"')
    format_run(r3, size_pt=7.5, color_rgb=(51, 65, 85))
    
    add_signatures_block(include_witnesses=True, include_thumbs=True)

    # Save DOCX
    docx_path = r"c:\Users\AC I\Desktop\MM\MM_Ride_Shivkumar_Driver_Document_Pack_English_Hindi.docx"
    doc.save(docx_path)
    print(f"DOCX created successfully: {docx_path}")
    return docx_path

if __name__ == "__main__":
    docx_path = create_document_pack()
    pdf_path = r"c:\Users\AC I\Desktop\MM\MM_Ride_Shivkumar_Driver_Document_Pack_English_Hindi.pdf"
    
    print("Converting DOCX to PDF...")
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        print(f"PDF created successfully via docx2pdf: {pdf_path}")
    except Exception as e:
        print(f"docx2pdf failed: {e}, trying win32com directly...")
        try:
            import win32com.client
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc_com = word.Documents.Open(docx_path)
            doc_com.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
            doc_com.Close()
            word.Quit()
            print(f"PDF created successfully via win32com: {pdf_path}")
        except Exception as e2:
            print(f"win32com failed: {e2}")
