import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_cell_border_box(cell, color="CCCCCC", sz="4"):
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_header_footer(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.27)  # A4
        section.page_height = Inches(11.69)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("MM RIDE | DRIVER DOCUMENT PACK (DL9SBH6153)")
        hrun.font.name = "Nirmala UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("MM Ride Confidential & Legal Driver Document Record  •  Shivkumar Shankarappa Nindi")
        frun.font.name = "Nirmala UI"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def format_run(run, font_name="Nirmala UI", size_pt=10, bold=False, italic=False, color_rgb=(30, 41, 59)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_bilingual_paragraph(doc, eng_text, hin_text, style_type="normal", space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if style_type == "h1":
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
        r1 = p.add_run(eng_text + "\n")
        format_run(r1, size_pt=14, bold=True, color_rgb=(15, 23, 42))
        r2 = p.add_run(hin_text)
        format_run(r2, size_pt=13, bold=True, color_rgb=(30, 58, 138))
    elif style_type == "h2":
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(eng_text + " / ")
        format_run(r1, size_pt=11.5, bold=True, color_rgb=(30, 58, 138))
        r2 = p.add_run(hin_text)
        format_run(r2, size_pt=11, bold=True, color_rgb=(30, 58, 138))
    elif style_type == "clause_title":
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(eng_text + "\n")
        format_run(r1, size_pt=10.5, bold=True, color_rgb=(15, 23, 42))
        r2 = p.add_run(hin_text)
        format_run(r2, size_pt=10, bold=True, color_rgb=(29, 78, 216))
    else: # normal clause body
        r1 = p.add_run(eng_text + "\n")
        format_run(r1, size_pt=9.5, color_rgb=(30, 41, 59))
        r2 = p.add_run(hin_text)
        format_run(r2, size_pt=9, italic=True, color_rgb=(71, 85, 105))
    return p

def add_doc_banner(doc, doc_num, eng_title, hin_title):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.77)
    set_cell_background(cell, "0F172A") # Deep Slate Navy
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    
    r0 = p.add_run(f"DOCUMENT {doc_num}\n")
    format_run(r0, size_pt=9, bold=True, color_rgb=(148, 163, 184))
    
    r1 = p.add_run(eng_title.upper() + "\n")
    format_run(r1, size_pt=13, bold=True, color_rgb=(255, 255, 255))
    
    r2 = p.add_run(hin_title)
    format_run(r2, size_pt=12, bold=True, color_rgb=(226, 232, 240))
    
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)

def add_signatures_block(doc, include_witnesses=True, include_thumbs=True):
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(12)
    p_head.paragraph_format.space_after = Pt(6)
    r = p_head.add_run("SIGNATURES & ACKNOWLEDGEMENT / हस्ताक्षर एवं अभिस्वीकृति")
    format_run(r, size_pt=10.5, bold=True, color_rgb=(15, 23, 42))
    
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(3.38), Inches(3.38)]
    
    # Row 0: Driver & Fleet Operator
    c_driver = table.cell(0, 0)
    c_op = table.cell(0, 1)
    
    c_driver.width = col_widths[0]
    c_op.width = col_widths[1]
    
    set_cell_border_box(c_driver, "CBD5E1")
    set_cell_border_box(c_op, "CBD5E1")
    set_cell_margins(c_driver, top=120, bottom=120, left=140, right=140)
    set_cell_margins(c_op, top=120, bottom=120, left=140, right=140)
    
    pd = c_driver.paragraphs[0]
    pd.paragraph_format.space_after = Pt(4)
    r = pd.add_run("DRIVER / ड्राइवर:\n")
    format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
    r = pd.add_run("Name: Shivkumar Shankarappa Nindi\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = pd.add_run("DL No: MH14 20100054576\n\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = pd.add_run("Signature: ___________________________\n\n")
    format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
    r = pd.add_run("Date: ______________  Place: ___________")
    format_run(r, size_pt=9, color_rgb=(71, 85, 105))
    
    if include_thumbs:
        pd.add_run("\n\nLeft Thumb: [        ]  Right Thumb: [        ]")
    
    po = c_op.paragraphs[0]
    po.paragraph_format.space_after = Pt(4)
    r = po.add_run("FLEET OPERATOR / फ्लीट ऑपरेटर:\n")
    format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
    r = po.add_run("Name: Mohamed Ibrahim\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = po.add_run("Role: Fleet Operator / Vehicle Custodian\n\n")
    format_run(r, size_pt=9, color_rgb=(30, 41, 59))
    r = po.add_run("Signature: ___________________________\n\n")
    format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
    r = po.add_run("Date: ______________  Place: ___________")
    format_run(r, size_pt=9, color_rgb=(71, 85, 105))

    if include_witnesses:
        c_w1 = table.cell(1, 0)
        c_w2 = table.cell(1, 1)
        
        c_w1.width = col_widths[0]
        c_w2.width = col_widths[1]
        
        set_cell_border_box(c_w1, "CBD5E1")
        set_cell_border_box(c_w2, "CBD5E1")
        set_cell_margins(c_w1, top=120, bottom=120, left=140, right=140)
        set_cell_margins(c_w2, top=120, bottom=120, left=140, right=140)
        
        pw1 = c_w1.paragraphs[0]
        pw1.paragraph_format.space_after = Pt(4)
        r = pw1.add_run("WITNESS 1 / गवाह 1:\n")
        format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
        r = pw1.add_run("Name: _____________________________\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw1.add_run("Address/Mobile: ____________________\n\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw1.add_run("Signature: _________________________\n")
        format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
        r = pw1.add_run("Date: _____________________________")
        format_run(r, size_pt=9, color_rgb=(71, 85, 105))
        
        pw2 = c_w2.paragraphs[0]
        pw2.paragraph_format.space_after = Pt(4)
        r = pw2.add_run("WITNESS 2 / गवाह 2:\n")
        format_run(r, size_pt=9.5, bold=True, color_rgb=(15, 23, 42))
        r = pw2.add_run("Name: _____________________________\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw2.add_run("Address/Mobile: ____________________\n\n")
        format_run(r, size_pt=9, color_rgb=(30, 41, 59))
        r = pw2.add_run("Signature: _________________________\n")
        format_run(r, size_pt=9, bold=True, color_rgb=(30, 41, 59))
        r = pw2.add_run("Date: _____________________________")
        format_run(r, size_pt=9, color_rgb=(71, 85, 105))

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_after = Pt(10)

print("Helper setup complete.")
